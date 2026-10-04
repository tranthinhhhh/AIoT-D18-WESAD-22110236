from __future__ import annotations

import json
import os
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "De_cuong_M7_Duong_Minh_Duy_23110083_DA_CHINH_SUA_biamau.docx"
OUTPUT = ROOT / "TranTienThinh_final.docx"
FIG = ROOT / "outputs" / "figures"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    node = OxmlElement("w:t")
    node.text = text
    run.extend([properties, node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def remove_after_cover(doc: Document) -> None:
    first_body_para = doc.paragraphs[15]._element
    body = doc._body._element
    remove = False
    for child in list(body):
        if child is first_body_para:
            remove = True
        if remove and child.tag != qn("w:sectPr"):
            body.remove(child)


def style_document(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color in [
        ("Title", 20, "17365D"), ("Heading 1", 16, "17365D"),
        ("Heading 2", 13, "1F4E79"), ("Heading 3", 11, "1F4E79"),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)
    if "Caption" in styles:
        cap = styles["Caption"]
        cap.font.name = "Times New Roman"
        cap._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        cap.font.size = Pt(9)
        cap.font.italic = True


def body(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.75)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.add_run(text)


def caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)


def picture(doc: Document, path: Path, caption_text: str, width: float = 6.25) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    caption(doc, caption_text)


def table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    hdr = tbl.rows[0]
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(value))
        r.bold = True
        r.font.size = Pt(9)
    for row in rows:
        cells = tbl.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            r.font.size = Pt(9)
    if widths:
        for row in tbl.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()


def page(doc: Document, title: str, number: str | None = None) -> None:
    doc.add_page_break()
    if number:
        doc.add_heading(f"{number}. {title}", level=1)
    else:
        doc.add_heading(title, level=1)


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def aggregate_condition(summary: dict) -> list[list[object]]:
    rows = []
    names = [entry["condition"] for entry in summary["runs"][0]["metrics"]["test"]["by_condition"]]
    for name in names:
        vals = []
        rmses = []
        n = None
        for run in summary["runs"]:
            item = next(x for x in run["metrics"]["test"]["by_condition"] if x["condition"] == name)
            vals.append(item["mae_macro_bpm"])
            rmses.append(item["rmse_macro_bpm"])
            n = item["n_windows"]
        rows.append([name, n, f"{np.mean(vals):.3f} ± {np.std(vals, ddof=1):.3f}", f"{np.mean(rmses):.3f} ± {np.std(rmses, ddof=1):.3f}"])
    return rows


def aggregate_subject(summary: dict) -> list[list[object]]:
    names = [entry["subject"] for entry in summary["runs"][0]["metrics"]["test"]["by_subject"]]
    rows = []
    for name in names:
        vals, rmses = [], []
        for run in summary["runs"]:
            item = next(x for x in run["metrics"]["test"]["by_subject"] if x["subject"] == name)
            vals.append(item["mae_bpm"])
            rmses.append(item["rmse_bpm"])
        rows.append([name, f"{np.mean(vals):.3f} ± {np.std(vals, ddof=1):.3f}", f"{np.mean(rmses):.3f} ± {np.std(rmses, ddof=1):.3f}"])
    return rows


def main() -> None:
    github_url = os.environ.get("GITHUB_URL", "").strip()
    shutil.copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    remove_after_cover(doc)
    cover = doc.paragraphs
    replacements = {
        4: "BÁO CÁO CUỐI KỲ",
        5: "MÔN TRÍ TUỆ NHÂN TẠO CHO IoT",
        6: "ĐỀ TÀI Đ18",
        7: "ƯỚC LƯỢNG NHỊP TIM TỪ TÍN HIỆU BVP\nCỔ TAY TRÊN WESAD BẰNG CNN 1-D,\nVỚI THAM CHIẾU ECG NGỰC VÀ ĐÁNH GIÁ\nTHEO ĐIỀU KIỆN THỰC NGHIỆM",
        8: "Wrist BVP Heart-Rate Estimation on WESAD Using a 1-D CNN\nwith Chest-ECG Reference and Condition-Wise Evaluation",
        9: "Giảng viên hướng dẫn: ThS. Hồ Nhựt Minh",
        10: "Sinh viên thực hiện: Trần Tiến Thịnh",
        11: "Mã số sinh viên: 22110236",
        12: "",
        13: "Năm học: 2026–2027",
        14: "Thành phố Hồ Chí Minh, năm 2026",
    }
    for index, value in replacements.items():
        cover[index].text = value
        cover[index].alignment = WD_ALIGN_PARAGRAPH.CENTER if index <= 8 or index >= 13 else WD_ALIGN_PARAGRAPH.LEFT
    style_document(doc)
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    for p in section.header.paragraphs:
        p.clear()
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_r = header_p.add_run("Đ18 • Trần Tiến Thịnh • 22110236")
    header_r.font.name = "Times New Roman"
    header_r.font.size = Pt(8)
    header_r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    for p in section.first_page_header.paragraphs:
        p.clear()
    for p in section.footer.paragraphs:
        p.clear()
    add_page_number(section.footer.paragraphs[0])
    for p in section.first_page_footer.paragraphs:
        p.clear()

    prep = load_json(ROOT / "data/processed/wesad_d18_v2/preprocessing_summary.json")
    manifest = load_json(ROOT / "data/processed/wesad_d18_v2/manifest.json")
    base = load_json(ROOT / "outputs/baselines/baseline_metrics.json")
    a0 = load_json(ROOT / "outputs/runs/A0_summary.json")
    a1 = load_json(ROOT / "outputs/runs/A1_summary.json")
    a3 = load_json(ROOT / "outputs/runs/A3_summary.json")
    bench = load_json(ROOT / "outputs/benchmark_cpu.json")
    analysis = load_json(ROOT / "outputs/analysis_summary.json")
    smoke = load_json(ROOT / "outputs/audit/smoke_test.json")
    qc = load_json(ROOT / "outputs/audit/qc_audit_summary.json")

    page(doc, "TÓM TẮT")
    body(doc, "Đồ án xây dựng và đánh giá một CNN 1-D gọn nhẹ để ước lượng nhịp tim trung bình từ duy nhất tín hiệu BVP cổ tay trên WESAD. ECG ngực chỉ được dùng để suy ra HR tham chiếu tự động; mã người và nhãn điều kiện không đi vào mô hình. Toàn bộ giao thức được khóa theo người: S2–S10 huấn luyện, S11/S13 validation và S14–S17 kiểm thử.")
    body(doc, f"Pipeline v2 tạo {prep['windows_before_qc']:,} cửa sổ, giữ {prep['windows_kept']:,} cửa sổ sau QC; tập train/validation/test lần lượt có {prep['windows_by_split']['train']:,}/{prep['windows_by_split']['validation']:,}/{prep['windows_by_split']['test']:,} cửa sổ. Trên cùng manifest, E0 đạt {base['methods']['E0_test']['mae_macro_bpm']:.3f} BPM và E1 đạt {base['methods']['E1_test']['mae_macro_bpm']:.3f} BPM. CNN A0 qua ba seed đạt {a0['aggregate_across_seeds']['mae_macro_bpm']['mean']:.3f} ± {a0['aggregate_across_seeds']['mae_macro_bpm']['std_sample']:.3f} BPM, giảm {a0['mae_macro_reduction_vs_fft_pct']:.1f}% so với FFT.")
    body(doc, f"A3 dùng MSE cho MAE_macro thấp nhất trong các cấu hình đã chạy: {a3['aggregate_across_seeds']['mae_macro_bpm']['mean']:.3f} ± {a3['aggregate_across_seeds']['mae_macro_bpm']['std_sample']:.3f} BPM. Mục tiêu giảm ít nhất 20% được đạt, nhưng mục tiêu tuyệt đối ≤7 BPM chưa đạt. Sai số tập trung ở trạng thái Stress và khác biệt rõ giữa người; do đó kết quả không được diễn giải như độ chính xác lâm sàng hoặc hiệu năng trên môi trường tự do.")
    table(doc, ["Thông tin", "Giá trị"], [["Dataset version", prep["dataset_version"]], ["Config hash", prep["config_hash"][:20] + "…"], ["Manifest hash", prep["manifest_hash"][:20] + "…"], ["Metric chính", "MAE_macro theo người"], ["Số seed", "42, 123, 999"]], [5, 11])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if github_url:
        r = p.add_run("Liên kết GitHub: ")
        r.bold = True
        add_hyperlink(p, github_url, github_url)
    else:
        r = p.add_run("Liên kết GitHub: CHƯA ĐƯỢC CẤU HÌNH — bản này chưa sẵn sàng nộp cho đến khi thay bằng URL repository thật.")
        r.bold = True
        r.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)

    page(doc, "MỤC LỤC")
    toc = [
        ("Tóm tắt", 2), ("Mục lục", 3), ("1. Đặt vấn đề", 4), ("2. Dữ liệu và ranh giới thông tin", 5),
        ("3. Tiền xử lý, nhãn tham chiếu và QC", 6), ("4. Phương pháp", 7), ("5. Thiết kế thực nghiệm", 8),
        ("6. Kết quả tổng thể", 9), ("7. Phân tầng theo người và điều kiện", 10),
        ("8. Phân tích đồng thuận", 11), ("9. Diễn biến theo thời gian và ca lỗi", 12),
        ("10. Ablation và động lực học huấn luyện", 13), ("11. Độ trễ và khả năng triển khai", 14),
        ("12. Thảo luận và giới hạn", 15), ("13. Tái lập và sản phẩm bàn giao", 16),
        ("14. Kết luận", 17), ("Tài liệu tham khảo", 18), ("Phụ lục", 19),
    ]
    table(doc, ["Mục", "Trang"], [[name, number] for name, number in toc], [14, 2])
    body(doc, "Danh mục viết tắt: BVP — Blood Volume Pulse; ECG — Electrocardiogram; HR — Heart Rate; BPM — beats per minute; QC — quality control; MAE — mean absolute error; RMSE — root mean squared error; FFT — Fast Fourier Transform.")

    page(doc, "ĐẶT VẤN ĐỀ", "1")
    doc.add_heading("1.1. Bối cảnh", level=2)
    body(doc, "BVP từ thiết bị đeo phản ánh biến thiên thể tích máu ngoại vi và có thể hỗ trợ ước lượng nhịp tim, nhưng chịu ảnh hưởng mạnh bởi vận động, tiếp xúc cảm biến và sinh lý từng người. WESAD cung cấp BVP cổ tay đồng bộ với ECG ngực trong một giao thức phòng thí nghiệm có Baseline, Stress và Amusement [1], [2].")
    doc.add_heading("1.2. Bài toán và câu hỏi", level=2)
    body(doc, "Với một cửa sổ BVP 8 giây sau hạ mẫu còn 256 điểm, mô hình hồi quy một giá trị HR trung bình. Câu hỏi chính là CNN 1-D có giảm MAE_macro so với FFT và đối chứng trung vị train trên cùng cửa sổ hay không; đồng thời sai số thay đổi ra sao giữa người và điều kiện.")
    doc.add_heading("1.3. Mục tiêu đăng ký trước", level=2)
    bullet(doc, "MAE_macro trên bốn người test không vượt quá 7 BPM.")
    bullet(doc, "Giảm ít nhất 20% MAE_macro so với E1 FFT.")
    bullet(doc, "Ba seed, checkpoint chọn theo validation; test không dùng để tuning.")
    bullet(doc, "Có kiểm tra nhãn ECG, phân tích lỗi, latency CPU và artifact tái lập.")
    body(doc, "Hai mục tiêu là tiêu chí kỹ thuật của đồ án, không phải ngưỡng lâm sàng. Nếu không đạt phải báo cáo trung thực, không chọn lại người/cửa sổ hoặc seed theo test.")

    page(doc, "DỮ LIỆU VÀ RANH GIỚI THÔNG TIN", "2")
    body(doc, "Dự án dùng 15 tệp người tham gia WESAD hiện có trong máy, không tải lại dataset. BVP 64 Hz là đầu vào duy nhất. ECG 700 Hz chỉ tạo HR tham chiếu; ACC/EDA/nhiệt độ không dùng. Ba nhãn điều kiện chỉ dùng lọc phạm vi và phân tầng báo cáo.")
    table(doc, ["Tập", "Đối tượng", "Số cửa sổ"], [["Train", "S2–S10", f"{prep['windows_by_split']['train']:,}"], ["Validation", "S11, S13", f"{prep['windows_by_split']['validation']:,}"], ["Test", "S14–S17", f"{prep['windows_by_split']['test']:,}"]], [3, 7, 4])
    body(doc, "Danh sách người được khóa trước huấn luyện. Kiểm tra tự động xác nhận window_id duy nhất, tensor và metadata thẳng hàng, không có người trùng giữa ba tập và mọi giá trị dùng để huấn luyện đều hữu hạn.")
    doc.add_heading("2.1. Phiên bản dữ liệu", level=2)
    table(doc, ["Thuộc tính", "Giá trị"], [["Phiên bản", prep["dataset_version"]], ["Định nghĩa cửa sổ", "[t, t+8 s), bước 2 s"], ["Số mẫu BVP", "256"], ["Config hash", prep["config_hash"]], ["Manifest hash", prep["manifest_hash"]]], [4, 12])

    page(doc, "TIỀN XỬ LÝ, NHÃN THAM CHIẾU VÀ QC", "3")
    body(doc, "BVP được hạ mẫu 64→32 Hz bằng resample_poly, lọc Butterworth SOS bậc 4 trong dải 0,5–8 Hz theo toàn bản ghi và chuẩn hóa z-score riêng từng cửa sổ. Đây là xử lý zero-phase ngoại tuyến, không được gọi là bộ lọc nhân quả trực tuyến.")
    body(doc, "ECG được lọc 5–25 Hz. Detector đỉnh R thích nghi chọn cực tính, dùng ngưỡng theo phân vị cục bộ, khoảng trơ và tinh chỉnh vị trí. Đây không phải triển khai Pan–Tompkins. Với các đỉnh nằm trong cùng cửa sổ, HR_ref = 60(n−1)/(t_n−t_1).")
    table(doc, ["Lý do chính", "Số cửa sổ"], [[key, f"{value:,}"] for key, value in prep["primary_reason_counts"].items()], [9, 5])
    body(doc, f"Tổng cộng loại {prep['windows_rejected']:,} cửa sổ: phần lớn do điều kiện ngoài phạm vi, không nên gọi chung là nhiễu. Bộ audit trực quan chọn {qc['accepted_windows_reviewed']} cửa sổ hợp lệ, đúng 5 cửa sổ cho mỗi tổ hợp 15 người × 3 điều kiện, và vẽ {qc['flagged_windows_plotted']} cửa sổ bị gắn cờ. Kiểm tra này đánh giá tính hợp lý hình dạng/nhãn tự động, không phải thẩm định y khoa.")
    picture(doc, FIG / "qc" / "S14_qc_15_windows.png", "Hình 1. Montage QC của S14; ECG chỉ minh họa nhãn và không vào CNN.", 5.55)

    page(doc, "PHƯƠNG PHÁP", "4")
    doc.add_heading("4.1. Đối chứng", level=2)
    body(doc, "E0 dự đoán trung vị HR_ref của train cho mọi cửa sổ. E1 nhân cửa sổ Hann, dùng NFFT=4096 và chọn đỉnh công suất lớn nhất trong miền 35/60–200/60 Hz. Zero-padding làm dày lưới tính toán lên 0,0078125 Hz nhưng không tạo thêm thông tin quan sát.")
    doc.add_heading("4.2. CNN 1-D", level=2)
    table(doc, ["Khối", "Cấu hình", "Đầu ra"], [["Conv1", "1→16, k=7, s=2, p=3 + BN + ReLU", "16×128"], ["Conv2", "16→32, k=5, s=2, p=2 + BN + ReLU", "32×64"], ["Conv3", "32→64, k=3, s=2, p=1 + BN + ReLU", "64×32"], ["Pooling", "AdaptiveAvgPool1d(1)", "64"], ["Head", "Dropout 0,2; Linear 64→32→1", "1 BPM"]], [3, 10, 3])
    body(doc, "A0/E2 dùng BVP đã lọc và z-score, Huber δ=5 BPM, BatchNorm và tổng 11.153 tham số học được. A1 bỏ z-score; A3 chỉ thay Huber bằng MSE. A2 bỏ BatchNorm nhưng chưa chạy để ưu tiên sản phẩm bắt buộc theo thời gian.")
    doc.add_heading("4.3. Ranh giới suy luận", level=2)
    body(doc, "Nhánh suy luận chỉ nhận BVP. ECG, HR_ref, subject ID và condition không xuất hiện trong tensor đầu vào. Nhờ đó chỉ số test đo khả năng tổng quát sang người chưa thấy theo split cố định.")

    page(doc, "THIẾT KẾ THỰC NGHIỆM", "5")
    table(doc, ["Hạng mục", "Cấu hình khóa"], [["Optimizer", "AdamW"], ["Learning rate / weight decay", "10⁻³ / 10⁻⁴"], ["Batch size", "64"], ["Epoch tối đa", "100"], ["Early stopping", "15 epoch không giảm validation MAE_macro"], ["Seed", "42, 123, 999"], ["Checkpoint", "Tốt nhất theo validation MAE_macro"], ["CPU threads", "1"], ["Metric chính", "MAE từng người rồi trung bình đều"]], [6, 10])
    body(doc, "Mỗi seed điều khiển khởi tạo, shuffle minibatch và bộ sinh ngẫu nhiên. Ba seed mô tả biến thiên huấn luyện trên cùng một split, không tương đương ba tập kiểm định độc lập. Sau khi checkpoint được khóa theo validation, test chỉ dùng đánh giá cuối.")
    body(doc, f"Smoke test xác nhận output {smoke['output_shape']}, {smoke['trainable_parameters']:,} tham số, loss/gradient hữu hạn và khả năng overfit tập con (loss {smoke['overfit_initial_loss']:.3f}→{smoke['overfit_final_loss']:.3f}).")
    doc.add_heading("5.1. Chỉ số", level=2)
    body(doc, "Báo cáo MAE_macro, RMSE_macro, MAE/RMSE pooled, bias, Pearson r, giới hạn đồng thuận mô tả và số lỗi dự đoán. So sánh với FFT dùng cùng kiểu tổng hợp MAE_macro. Các cửa sổ chồng lấn không được xem là quan sát độc lập để suy luận thống kê.")

    page(doc, "KẾT QUẢ TỔNG THỂ", "6")
    overall = pd.read_csv(ROOT / "outputs/tables/overall_results_v2.csv")
    rows = []
    for _, item in overall.iterrows():
        std = "—" if pd.isna(item.mae_macro_std_bpm) else f"{item.mae_macro_std_bpm:.3f}"
        rows.append([item.method, int(item.seeds), f"{item.mae_macro_mean_bpm:.3f}", std, f"{item.pooled_rmse_mean_bpm:.3f}", f"{item.pearson_mean:.3f}"])
    table(doc, ["Phương pháp", "Seed", "MAE_macro", "SD", "RMSE pooled", "Pearson r"], rows, [3, 2, 3, 2, 3, 3])
    picture(doc, FIG / "overall_mae_macro.png", "Hình 2. MAE_macro test trên cùng manifest v2; error bar là SD mẫu qua seed.", 5.8)
    body(doc, f"A0 đạt {a0['aggregate_across_seeds']['mae_macro_bpm']['mean']:.3f} ± {a0['aggregate_across_seeds']['mae_macro_bpm']['std_sample']:.3f} BPM, giảm {a0['mae_macro_reduction_vs_fft_pct']:.1f}% so với E1. A3 đạt {a3['aggregate_across_seeds']['mae_macro_bpm']['mean']:.3f} ± {a3['aggregate_across_seeds']['mae_macro_bpm']['std_sample']:.3f} BPM và là cấu hình thấp nhất trong các run đã thực hiện. Tuy nhiên, cả A0 và A3 đều cao hơn mục tiêu 7 BPM.")

    page(doc, "PHÂN TẦNG THEO NGƯỜI VÀ ĐIỀU KIỆN", "7")
    doc.add_heading("7.1. Theo người — A0 qua ba seed", level=2)
    table(doc, ["Người test", "MAE (BPM)", "RMSE (BPM)"], aggregate_subject(a0), [4, 6, 6])
    doc.add_heading("7.2. Theo điều kiện — A0 qua ba seed", level=2)
    table(doc, ["Điều kiện", "Số cửa sổ", "MAE_macro", "RMSE_macro"], aggregate_condition(a0), [4, 4, 5, 5])
    picture(doc, FIG / "error_by_subject_condition.png", "Hình 3. Sai số A0 seed 42 theo người và điều kiện.", 5.9)
    body(doc, "Baseline và Amusement có sai số thấp, trong khi Stress chi phối lỗi. Đây là mô tả trên split hiện tại; không được suy diễn Stress gây ra sai số vì HR, chuyển động, subject và chất lượng tín hiệu đồng biến.")

    page(doc, "PHÂN TÍCH ĐỒNG THUẬN", "8")
    picture(doc, FIG / "scatter_identity.png", "Hình 4. HR tham chiếu–dự đoán A0 seed 42 với đường đồng nhất.", 5.65)
    picture(doc, FIG / "bland_altman_v2.png", "Hình 5. Bland–Altman mô tả A0 seed 42; các cửa sổ lặp/chồng lấn không độc lập.", 5.65)
    body(doc, f"A0 qua ba seed có Pearson r pooled {a0['aggregate_across_seeds']['pooled_pearson_r']['mean']:.3f} ± {a0['aggregate_across_seeds']['pooled_pearson_r']['std_sample']:.3f} và bias {a0['aggregate_across_seeds']['pooled_bias_bpm']['mean']:.3f} ± {a0['aggregate_across_seeds']['pooled_bias_bpm']['std_sample']:.3f} BPM. Bias âm phản ánh xu hướng đánh giá thấp ở vùng HR cao, rõ nhất trong Stress.")

    page(doc, "DIỄN BIẾN THEO THỜI GIAN VÀ CA LỖI", "9")
    picture(doc, FIG / "timeline_s14_fixed.png", "Hình 6. HR theo thời gian của S14, seed 42 được cố định trước phân tích test.", 6.1)
    failure = analysis["failure_case"]
    picture(doc, FIG / "failure_case_validation_fixed.png", "Hình 7. Ca lỗi điển hình được chọn trên validation, không chọn từ test.", 6.1)
    body(doc, f"Ca lỗi là {failure['window_id']} ({failure['condition']}): HR_ref {failure['hr_ref_bpm']:.1f} BPM, A0 dự đoán {failure['hr_pred_bpm']:.1f} BPM, sai số {failure['error_bpm']:.1f} BPM. Hình kèm BVP, phổ FFT, ECG và đỉnh R giúp phân biệt lỗi mô hình với bất thường nhãn; nhãn vẫn chỉ là tham chiếu tự động.")

    page(doc, "ABLATION VÀ ĐỘNG LỰC HỌC HUẤN LUYỆN", "10")
    table(doc, ["Biến thể", "Thay đổi duy nhất", "MAE_macro (BPM)", "Giảm so E1"], [["A0", "Chuẩn hóa + Huber + BN", f"{a0['aggregate_across_seeds']['mae_macro_bpm']['mean']:.3f} ± {a0['aggregate_across_seeds']['mae_macro_bpm']['std_sample']:.3f}", f"{a0['mae_macro_reduction_vs_fft_pct']:.1f}%"], ["A1", "Bỏ z-score", f"{a1['aggregate_across_seeds']['mae_macro_bpm']['mean']:.3f} ± {a1['aggregate_across_seeds']['mae_macro_bpm']['std_sample']:.3f}", f"{a1['mae_macro_reduction_vs_fft_pct']:.1f}%"], ["A2", "Bỏ BatchNorm", "Chưa chạy", "—"], ["A3", "MSE thay Huber", f"{a3['aggregate_across_seeds']['mae_macro_bpm']['mean']:.3f} ± {a3['aggregate_across_seeds']['mae_macro_bpm']['std_sample']:.3f}", f"{a3['mae_macro_reduction_vs_fft_pct']:.1f}%"]], [3, 6, 5, 4])
    picture(doc, FIG / "learning_curves_v2.png", "Hình 8. Đường học train và validation của các run đã hoàn tất.", 6.1)
    body(doc, "A1 kém A0, cho thấy chuẩn hóa theo cửa sổ hữu ích trong cấu hình này. A3 tốt hơn A0 ở cả trung bình ba seed, trái với kỳ vọng Huber luôn bền vững hơn; đây là kết quả thực nghiệm cần được báo cáo, nhưng chưa đủ để khái quát ngoài split. A2 được ghi rõ chưa chạy thay vì gọi là hoàn thành.")

    page(doc, "ĐỘ TRỄ VÀ KHẢ NĂNG TRIỂN KHAI", "11")
    table(doc, ["Phép đo CPU, batch=1", "Kết quả"], [["Warm-up / lượt đo mạng", f"{bench['warmup_iterations_network']} / {bench['measured_iterations_network']}"], ["Median mạng", f"{bench['network_only']['median_ms']:.4f} ms"], ["P95 mạng", f"{bench['network_only']['p95_ms']:.4f} ms"], ["Tiền xử lý toàn bản ghi S14", f"{bench['preprocessing_offline_whole_record']['median_ms']:.3f} ms"], ["Ước lượng cộng dồn median", f"{bench['end_to_end_additive_offline_estimate_ms']['median']:.4f} ms"], ["Tham số", f"{bench['trainable_parameters']:,}"], ["Checkpoint", f"{bench['checkpoint_bytes']:,} byte"], ["CPU / PyTorch", f"{bench['cpu']} / {bench['torch']}"]], [8, 8])
    body(doc, "Độ trễ mạng rất nhỏ so với chu kỳ cập nhật 2 giây. Tuy nhiên, cửa sổ cần 8 giây dữ liệu và phép lọc hiện tại là offline toàn bản ghi. Vì vậy con số cộng dồn chỉ là ước lượng amortized, không chứng minh hệ thống streaming có end-to-end latency dưới 1 ms.")
    body(doc, "Kết quả đủ để xem CNN là ứng viên cho gateway/edge CPU. Muốn khẳng định chạy trên MCU phải thiết kế bộ lọc nhân quả, đo RAM/flash và benchmark trên đúng phần cứng đích.")

    page(doc, "THẢO LUẬN VÀ GIỚI HẠN", "12")
    doc.add_heading("12.1. Điều đã chứng minh", level=2)
    bullet(doc, "Trên cùng manifest v2, CNN giảm MAE_macro hơn 20% so với FFT đã cài đặt.")
    bullet(doc, "Kết quả ba seed, log, checkpoint và prediction được giữ đầy đủ; không chọn seed tốt nhất theo test.")
    bullet(doc, "Metric chính là trung bình đều theo người, tránh để người có nhiều cửa sổ chi phối.")
    doc.add_heading("12.2. Điều chưa chứng minh", level=2)
    bullet(doc, "Không đạt mục tiêu ≤7 BPM; không có cơ sở gọi mô hình đạt độ chính xác lâm sàng.")
    bullet(doc, "Chỉ bốn người test, dữ liệu phòng thí nghiệm và các cửa sổ chồng lấn; độ bất định theo quần thể còn lớn.")
    bullet(doc, "ECG label là tham chiếu tự động và detector chưa được chuyên gia y khoa thẩm định.")
    bullet(doc, "A2 chưa chạy; không tuyên bố ablation đầy đủ 4×3 run.")
    bullet(doc, "Không so sánh với mọi thuật toán HR hiện đại; E1 là FFT đơn giản đã đặc tả.")
    doc.add_heading("12.3. Hướng cải thiện", level=2)
    body(doc, "Ưu tiên tiếp theo là mô hình có vùng tiếp nhận dài hơn hoặc đặc trưng nhịp, hiệu chỉnh nhãn ECG theo mẫu phân tầng, đánh giá LOSO có validation nội bộ, và pipeline streaming nhân quả. Mọi thay đổi phải tạo protocol mới thay vì tối ưu trực tiếp trên test hiện tại.")

    page(doc, "TÁI LẬP VÀ SẢN PHẨM BÀN GIAO", "13")
    table(doc, ["Artifact", "Vai trò"], [["configs/final.yaml", "Cấu hình duy nhất của protocol v2"], ["data/processed/wesad_d18_v2/manifest.json", "Nguồn, hash, SOS, phiên bản thư viện"], ["src/", "Preprocess, baseline, train/resume, metrics, analysis, benchmark, QC"], ["outputs/runs/", "best/last checkpoint, log, prediction và metric mỗi seed"], ["TranTienThinh_final.ipynb", "Notebook đánh giá đã chạy hết"], ["TranTienThinh_final.pptx", "Slide thuyết trình 5–7 phút"], ["submission/", "Gói nộp không chứa raw WESAD/venv/cache"]], [7, 11])
    body(doc, "Mỗi run chứa run_hash, config_hash và manifest_hash. Lệnh train trả lại metric cũ nếu run đã hoàn tất và hash khớp; resume chỉ chấp nhận last checkpoint tương thích. Cơ chế này ngăn huấn luyện trùng và tái sử dụng sai protocol.")
    body(doc, "Raw WESAD không được đóng gói lại. Người tái lập phải tự đặt dữ liệu theo README và tuân thủ điều kiện sử dụng. Notebook cuối kỳ mặc định đánh giá checkpoint; lệnh full train đặt riêng để tránh vô tình huấn luyện lại.")
    p = doc.add_paragraph()
    if github_url:
        r = p.add_run("Repository mã nguồn và artifact: ")
        r.bold = True
        add_hyperlink(p, github_url, github_url)
    else:
        r = p.add_run("Trạng thái GitHub: CHƯA CÓ REMOTE ĐƯỢC XÁC MINH. Cần thay dòng này bằng liên kết thật trước khi nộp.")
        r.bold = True
        r.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)

    page(doc, "KẾT LUẬN", "14")
    body(doc, "Đồ án đã hoàn thiện một pipeline tái lập BVP→HR trên WESAD với split theo người, QC có dấu vết, ba seed cho E2/A0 và A3, cùng đối chứng E0/E1 trên cùng manifest. A0 đạt 8,867 ± 0,253 BPM và giảm 35,0% so với FFT. A3 đạt 8,280 ± 0,328 BPM, tốt nhất trong các cấu hình đã chạy.")
    body(doc, "Mục tiêu giảm tương đối được đạt nhưng mục tiêu tuyệt đối 7 BPM không đạt. Phân tích cho thấy hạn chế tập trung ở Stress và tính không đồng nhất giữa người. Vì pipeline dùng lọc offline và nhãn ECG tự động, kết luận phù hợp nhất là: CNN gọn nhẹ có lợi thế rõ so với FFT đã cài đặt trên split cố định, nhưng chưa đủ bằng chứng cho sử dụng lâm sàng hoặc triển khai MCU trực tuyến.")
    body(doc, "Giá trị chính của sản phẩm nằm ở tính minh bạch: mọi con số trong báo cáo dẫn đến artifact cụ thể, ca lỗi chọn từ validation, test không dùng để tuning và phần chưa chạy được ghi đúng trạng thái.")

    page(doc, "TÀI LIỆU THAM KHẢO")
    refs = [
        "[1] P. Schmidt, A. Reiss, R. Dürichen, C. Marberger, K. Van Laerhoven, “Introducing WESAD,” ICMI, 2018. DOI: 10.1145/3242969.3242985.",
        "[2] P. Schmidt và A. Reiss, WESAD, UCI Machine Learning Repository, 2018. DOI: 10.24432/C57K5T.",
        "[3] A. Reiss et al., “Deep PPG,” Sensors, 19(14), 3079, 2019. DOI: 10.3390/s19143079.",
        "[4] P. H. Charlton et al., “The 2023 wearable photoplethysmography roadmap,” Physiological Measurement, 44(11), 111001, 2023.",
        "[5] J. Pan và W. J. Tompkins, “A Real-Time QRS Detection Algorithm,” IEEE TBME, 32(3), 230–236, 1985.",
        "[6] SciPy Community, SciPy signal processing documentation: resample_poly, sosfiltfilt và spectral analysis.",
        "[7] PyTorch Contributors, PyTorch documentation: torch.nn.HuberLoss và torch.optim.AdamW.",
        "[8] J. M. Bland và D. G. Altman, “Agreement between methods of measurement with multiple observations per individual,” Journal of Biopharmaceutical Statistics, 17(4), 571–582, 2007.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.add_run(ref)

    page(doc, "PHỤ LỤC A — LỆNH TÁI LẬP")
    commands = [
        "python -m src.pipeline preprocess",
        "python -m src.pipeline baselines",
        "python -m src.pipeline smoke",
        "python -m src.pipeline train --variant A0 --seed all",
        "python -m src.pipeline train --variant A3 --seed all",
        "python -m src.pipeline evaluate --variant A1 --seed all",
        "python -m src.analysis",
        "python -m src.benchmark",
        "python -m src.qc_audit",
    ]
    for command in commands:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1)
        r = p.add_run(command)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
    body(doc, "Hash đang dùng: config " + prep["config_hash"] + "; manifest " + prep["manifest_hash"] + ".")
    doc.add_heading("Phụ lục B — Đối chiếu kết quả theo seed", level=2)
    seed_rows = []
    for variant, summary in [("A0", a0), ("A1", a1), ("A3", a3)]:
        for seed, value in zip(summary["seeds"], summary["aggregate_across_seeds"]["mae_macro_bpm"]["values"]):
            seed_rows.append([variant, seed, f"{value:.6f}"])
    table(doc, ["Biến thể", "Seed", "Test MAE_macro (BPM)"], seed_rows, [4, 4, 8])

    props = doc.core_properties
    props.title = "Báo cáo cuối kỳ Đ18 — Ước lượng nhịp tim từ BVP cổ tay trên WESAD"
    props.author = "Trần Tiến Thịnh — 22110236"
    props.subject = "Trí tuệ nhân tạo cho IoT"
    props.keywords = "WESAD, BVP, ECG, CNN 1-D, heart rate, AIoT"
    doc.save(OUTPUT)
    print(OUTPUT.name)


if __name__ == "__main__":
    main()
